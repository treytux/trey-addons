###############################################################################
# For copyright and license notices, see __manifest__.py file in root directory
###############################################################################
import base64
import io
import os

from odoo import exceptions
from odoo.tests import common
from odoo.tools import load_language
from openpyxl import load_workbook


class TestSaleOrderLinesByRef(common.TransactionCase):

    def setUp(self):
        super().setUp()
        self.partner = self.env['res.partner'].create({
            'name': 'Test partner',
            'is_company': True,
        })
        self.product_1 = self.env['product.product'].create({
            'type': 'service',
            'sale_ok': True,
            'name': 'One',
            'list_price': 10,
            'standard_price': 3,
            'default_code': 'R1',
        })
        self.product_2 = self.env['product.product'].create({
            'type': 'service',
            'sale_ok': True,
            'name': 'Two',
            'list_price': 20,
            'default_code': 'R2',
            'barcode': '0123456789104',
        })

    def test_sale_order(self):
        sale = self.env['sale.order'].create({
            'partner_id': self.partner.id,
        })
        wizard = self.env['sale.order.lines_by_ref'].create({
            'sale_id': sale.id,
            'references': 'R1/10/30.10',
        })
        wizard.action_simulate()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_create()
        line = sale.order_line
        self.assertEqual(len(line), 1)
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 10)
        self.assertEqual(line.price_unit, 30.10)
        sale.order_line.unlink()
        self.assertEqual(len(sale.order_line), 0)
        wizard = self.env['sale.order.lines_by_ref'].create({
            'sale_id': sale.id,
            'references': '\n'.join(['0123456789104/10/20']),
        })
        wizard.action_simulate()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_create()
        line = sale.order_line
        self.assertEqual(len(line), 1)
        self.assertEqual(line.product_id, self.product_2)
        self.assertEqual(line.product_uom_qty, 10)
        self.assertEqual(line.price_unit, 20)
        sale.order_line.unlink()
        self.assertEqual(len(sale.order_line), 0)
        wizard = self.env['sale.order.lines_by_ref'].create({
            'sale_id': sale.id,
            'references': '\n'.join(['R1', 'R2']),
        })
        wizard.action_simulate()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_create()
        lines = sale.order_line
        self.assertEqual(len(lines), 2)
        self.assertEqual(lines[0].product_id, self.product_1)
        self.assertEqual(lines[0].product_uom_qty, 1)
        self.assertEqual(lines[0].price_unit, 10)
        self.assertEqual(lines[1].product_id, self.product_2)
        self.assertEqual(lines[1].product_uom_qty, 1)
        self.assertEqual(lines[1].price_unit, 20)

    def test_sale_order_with_errors(self):
        sale = self.env['sale.order'].create({
            'partner_id': self.partner.id,
        })
        wizard = self.env['sale.order.lines_by_ref'].create({
            'sale_id': sale.id,
            'references': 'RXXX/10/30.10',
        })
        wizard.action_simulate()
        self.assertEqual(len(wizard.line_ids), 1)
        self.assertEqual('Ref not exists', wizard.line_ids.name)
        wizard.action_create()
        self.assertEqual(len(sale.order_line), 0)

    def test_change_glue_char(self):
        set_param = self.env['ir.config_parameter'].sudo().set_param
        set_param('sale_order_lines_by_ref.glue', ',')
        sale = self.env['sale.order'].create({
            'partner_id': self.partner.id,
        })
        wizard = self.env['sale.order.lines_by_ref'].create({
            'sale_id': sale.id,
            'references': 'R1,10,30.10',
        })
        wizard.action_simulate()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_create()
        self.assertEqual(len(sale.order_line), 1)

    def test_ref_not_for_sale(self):
        sale = self.env['sale.order'].create({
            'partner_id': self.partner.id,
        })
        self.product_1.sale_ok = False
        wizard = self.env['sale.order.lines_by_ref'].create({
            'sale_id': sale.id,
            'references': 'R1',
        })
        wizard.action_simulate()
        self.assertEqual(len(wizard.line_ids), 1)
        self.assertEqual('Ref not exists', wizard.line_ids.name)
        wizard.action_create()
        self.assertEqual(len(sale.order_line), 0)

    def get_sample(self, fname):
        return os.path.join(os.path.dirname(__file__), fname)

    def test_import_xlsx_lines_ok(self):
        fname = self.get_sample('sample_ok.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        sale = self.env['sale.order'].create({
            'partner_id': self.partner.id,
        })
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, 8)

    def test_import_xlsx_lines_ok_spanish(self):
        fname = self.get_sample('sample_ok_spanish.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        sale = self.env['sale.order'].create({
            'partner_id': self.partner.id,
        })
        wizard = self.env['sale.order.lines_by_ref.xlsx'].with_context(
            lang='es_ES').create({
                'sale_id': sale.id,
                'xlsx_file': file,
            })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, 8)

    def test_action_provide_template_xlsx_spanish(self):
        load_language(self.env.cr, 'es_ES')
        sale = self.env['sale.order'].create({
            'partner_id': self.partner.id,
        })
        wizard = self.env['sale.order.lines_by_ref.xlsx'].with_context(
            lang='es_ES'
        ).create({
            'sale_id': sale.id,
        })
        result = wizard.action_provide_template_xlsx()
        self.assertEqual(result['type'], 'ir.actions.act_url')
        self.assertIn('/web/content/', result['url'])
        attachment = self.env['ir.attachment'].search([
            ('res_model', '=', 'sale.order.lines_by_ref.xlsx'),
            ('res_id', '=', wizard.id),
        ], limit=1, order='id desc')
        self.assertTrue(attachment)
        self.assertEqual(attachment.name, wizard._template_filename())
        data = base64.b64decode(attachment.datas)
        workbook = load_workbook(io.BytesIO(data))
        self.assertEqual(workbook.sheetnames, [
            wizard._template_sheet_name(),
            wizard._template_instruction_sheet_name(),
        ])
        main_sheet = workbook[wizard._template_sheet_name()]
        instruction_sheet = workbook[wizard._template_instruction_sheet_name()]
        expected_headers = [label for _, label in wizard._xlsx_columns()]
        for col, expected in enumerate(expected_headers, start=1):
            self.assertEqual(
                main_sheet.cell(row=1, column=col).value, expected)
        for row, expected in enumerate(
                wizard._template_instructions(), start=1):
            self.assertEqual(
                instruction_sheet.cell(row=row, column=1).value, expected)

    def test_import_xlsx_lines_error_ref_not_exists(self):
        fname = self.get_sample('sample_ref_not_exists.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        sale = self.env['sale.order'].create({
            'partner_id': self.partner.id,
        })
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 1)
        self.assertEqual('Ref not exists', wizard.line_ids.name)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 0)

    def test_import_xlsx_lines_ref_not_exists_with_fallback_product(self):
        fname = self.get_sample('sample_ref_not_exists.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        sale = self.env['sale.order'].create({
            'partner_id': self.partner.id,
        })
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
            'fallback_product_id': self.product_2.id,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 1)
        self.assertEqual('warning', wizard.line_ids.type)
        self.assertEqual(
            'Ref not exists. The fallback product will be used.',
            wizard.line_ids.name)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_2)
        self.assertEqual(line.name, 'REF_NOT_EXISTS')
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, 8)

    def test_import_xlsx_lines_error_sale_ok_false(self):
        self.product_1.sale_ok = False
        self.assertFalse(self.product_1.sale_ok)
        fname = self.get_sample('sample_sale_ok_false.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        sale = self.env['sale.order'].create({
            'partner_id': self.partner.id,
        })
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 1)
        self.assertEqual('Ref not exists', wizard.line_ids.name)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 0)

    def test_import_xlsx_lines_error_columns_number(self):
        fname = self.get_sample('sample_error_columns_number.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        sale = self.env['sale.order'].create({
            'partner_id': self.partner.id,
        })
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        with self.assertRaises(exceptions.ValidationError) as result:
            wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(
            result.exception.args[0], 'The file must have 3 columns')
        with self.assertRaises(exceptions.ValidationError) as result:
            wizard.action_import_lines_from_xlsx()
        self.assertEqual(
            result.exception.args[0], 'The file must have 3 columns')

    def _customer_info(self):
        modules = self.env['ir.module.module'].search([
            ('name', '=', 'product_supplierinfo_for_customer_sale'),
            ('state', '=', 'installed'),
        ])
        return modules and True or False

    def _pricelist_formula_margin(self):
        modules = self.env['ir.module.module'].search([
            ('name', '=', 'product_pricelist_formula_margin'),
            ('state', '=', 'installed'),
        ])
        return modules and True or False

    def test_import_xlsx_partner_shipping_customerinfo(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        fname = self.get_sample('sample_ok.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        partner = self.env.ref('base.res_partner_12')
        partner_shipping = self.env['res.partner'].create({
            'name': 'Partner shipping test',
        })
        self.env['product.customerinfo'].create({
            'product_id': self.product_1.id,
            'partner_id': partner_shipping.id,
            'product_code': '2323',
        })
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
            'partner_shipping_id': partner_shipping.id,
        })
        customerinfo_obj = self.env['product.customerinfo']
        partner_customerinfo = customerinfo_obj.search([
            ('product_code', '=', '2323'),
            ('partner_id', '=', partner.id),
        ])
        partner_shipping_customerinfo = customerinfo_obj.search([
            ('product_code', '=', '2323'),
            ('partner_id', '=', partner_shipping.id),
        ])
        self.assertEqual(len(partner_customerinfo), 0)
        self.assertEqual(len(partner_shipping_customerinfo), 1)
        self.assertEqual(sale.partner_id, partner)
        self.assertEqual(sale.partner_shipping_id, partner_shipping)
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, 8)

    def test_import_xlsx_partner_shipping_customerinfo_02(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        fname = self.get_sample('sample_ok_ends_zero.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        partner = self.env.ref('base.res_partner_12')
        partner_shipping = self.env['res.partner'].create({
            'name': 'Partner shipping test',
        })
        self.env['product.customerinfo'].create({
            'product_id': self.product_1.id,
            'partner_id': partner_shipping.id,
            'product_code': '9128812',
        })
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
            'partner_shipping_id': partner_shipping.id,
        })
        customerinfo_obj = self.env['product.customerinfo']
        partner_customerinfo = customerinfo_obj.search([
            ('product_code', '=', '9128812'),
            ('partner_id', '=', partner.id),
        ])
        partner_shipping_customerinfo = customerinfo_obj.search([
            ('product_code', '=', '9128812'),
            ('partner_id', '=', partner_shipping.id),
        ])
        self.assertEqual(len(partner_customerinfo), 0)
        self.assertEqual(len(partner_shipping_customerinfo), 1)
        self.assertEqual(sale.partner_id, partner)
        self.assertEqual(sale.partner_shipping_id, partner_shipping)
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 100)
        self.assertEqual(line.price_unit, 18)

    def test_sale_no_price_xlsx_product_customerinfo_01(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        fname = self.get_sample('sample_ok_customerinfo_1.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        partner = self.env.ref('base.res_partner_12')
        partner_shipping = self.env['res.partner'].create({
            'name': 'Partner shipping test',
        })
        self.env['product.customerinfo'].create({
            'product_id': self.product_1.id,
            'partner_id': partner_shipping.id,
            'product_code': 'R2323',
            'price': 15,
        })
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
            'partner_shipping_id': partner_shipping.id,
        })
        customerinfo_obj = self.env['product.customerinfo']
        partner_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner.id),
        ])
        partner_shipping_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner_shipping.id),
        ])
        self.assertEqual(len(partner_customerinfo), 0)
        self.assertEqual(len(partner_shipping_customerinfo), 1)
        self.assertEqual(sale.partner_id, partner)
        self.assertEqual(sale.partner_shipping_id, partner_shipping)
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, 15)

    def test_sale_no_price_xlsx_product_customerinfo_02(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        fname = self.get_sample('sample_ok_customerinfo_1.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        partner = self.env.ref('base.res_partner_12')
        partner_shipping = self.env['res.partner'].create({
            'name': 'Partner shipping test',
        })
        self.env['product.customerinfo'].create({
            'product_id': self.product_1.id,
            'partner_id': partner.id,
            'product_code': 'R2323',
            'price': 15,
        })
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
        })
        customerinfo_obj = self.env['product.customerinfo']
        partner_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner.id),
        ])
        partner_shipping_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner_shipping.id),
        ])
        self.assertEqual(len(partner_customerinfo), 1)
        self.assertEqual(len(partner_shipping_customerinfo), 0)
        self.assertEqual(sale.partner_id, partner)
        self.assertEqual(sale.partner_shipping_id, partner)
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, 15)

    def test_sale_no_price_xlsx_product_list_price_01(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        fname = self.get_sample('sample_ok_customerinfo_2.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        partner = self.env.ref('base.res_partner_12')
        partner_shipping = self.env['res.partner'].create({
            'name': 'Partner shipping test',
        })
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
            'partner_shipping_id': partner_shipping.id,
        })
        self.assertEqual(sale.partner_id, partner)
        self.assertEqual(sale.partner_shipping_id, partner_shipping)
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        self.assertEqual(len(sale.order_line), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, self.product_1.list_price)

    def test_sale_no_price_xlsx_product_list_price_02(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        fname = self.get_sample('sample_ok_customerinfo_2.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        partner = self.env.ref('base.res_partner_12')
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
        })
        self.assertEqual(sale.partner_id, partner)
        self.assertEqual(sale.partner_shipping_id, partner)
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        self.assertEqual(len(sale.order_line), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, self.product_1.list_price)

    def test_sale_no_price_xlsx_product_price_by_pricelist_01(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        fname = self.get_sample('sample_ok_customerinfo_2.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        partner = self.env.ref('base.res_partner_12')
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
        })
        self.pricelist_item = self.env['product.pricelist.item'].create({
            'applied_on': '1_product',
            'base': 'partner',
            'name': 'Test pricelist item',
            'pricelist_id': sale.pricelist_id.id,
            'compute_price': 'fixed',
            'fixed_price': 6,
            'product_id': self.product_1.id,
        })
        self.assertEqual(sale.partner_id, partner)
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        self.assertEqual(len(sale.order_line), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, self.pricelist_item.fixed_price)
        self.assertEqual(
            line.price_unit, self.product_1.with_context(
                pricelist=sale.pricelist_id.id).price)

    def test_sale_order_lines_customer(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        partner = self.env.ref('base.res_partner_12')
        product = self.env['product.product'].create({
            'name': 'Product Test',
            'type': 'consu',
            'product_category': self.env.ref('product.product_category_all'),
        })
        self.env['product.customerinfo'].create({
            'product_id': product.id,
            'partner_id': partner.id,
            'product_code': '2323',
        })
        sale_order = self.env['sale.order'].create({
            'partner_id': partner.id,
        })
        products_customer_code = self.env['product.customerinfo'].search([
            ('product_code', '=', '2323'),
            ('partner_id', '=', partner.id),
        ])
        wizard = self.env['sale.order.lines_by_ref'].create({
            'sale_id': sale_order.id,
            'references': '2323',
        })
        wizard.action_simulate()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_create()
        self.assertEqual(len(sale_order.order_line), 1)
        self.assertEqual(len(products_customer_code), 1)
        self.assertEqual(len(products_customer_code.product_id), 1)
        self.assertEqual(len(products_customer_code.product_tmpl_id), 0)

    def test_sale_order_lines_same_customer_code_error(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        partner = self.env.ref('base.res_partner_12')
        product = self.env['product.product'].create({
            'name': 'Product Test',
            'type': 'consu',
            'product_category': self.env.ref('product.product_category_all'),
        })
        self.env['product.customerinfo'].create({
            'product_id': product.id,
            'partner_id': partner.id,
            'product_code': '2323',
        })
        self.env['product.customerinfo'].create({
            'product_id': product.id,
            'partner_id': partner.id,
            'product_code': '2323',
        })
        sale_order = self.env['sale.order'].create({
            'partner_id': partner.id,
        })
        products_customer_code = self.env['product.customerinfo'].search([
            ('product_code', '=', '2323'),
            ('partner_id', '=', partner.id),
        ])
        wizard = self.env['sale.order.lines_by_ref'].create({
            'sale_id': sale_order.id,
            'references': '2323',
        })
        wizard.action_simulate()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_create()
        self.assertEqual(len(sale_order.order_line), 1)
        self.assertEqual(len(products_customer_code), 2)
        self.assertEqual(len(products_customer_code[0].product_id), 1)
        self.assertEqual(len(products_customer_code[0].product_tmpl_id), 0)

    def test_sale_order_lines_customer_product_template(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        partner = self.env.ref('base.res_partner_12')
        product = self.env['product.product'].create({
            'name': 'Product Test',
            'type': 'consu',
            'product_category': self.env.ref('product.product_category_all'),
        })
        self.env['product.customerinfo'].create({
            'product_tmpl_id': product.product_tmpl_id.id,
            'partner_id': partner.id,
            'product_code': '2323',
        })
        sale_order = self.env['sale.order'].create({
            'partner_id': partner.id,
        })
        products_customer_code = self.env['product.customerinfo'].search([
            ('product_code', '=', '2323'),
            ('partner_id', '=', partner.id),
        ])
        wizard = self.env['sale.order.lines_by_ref'].create({
            'sale_id': sale_order.id,
            'references': '2323',
        })
        wizard.action_simulate()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_create()
        self.assertEqual(len(sale_order.order_line), 1)
        self.assertEqual(len(products_customer_code), 1)
        self.assertEqual(len(products_customer_code[0].product_id), 0)
        self.assertEqual(len(products_customer_code[0].product_tmpl_id), 1)

    def test_sale_order_lines_customer_product_template_variant_errors(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        partner = self.env.ref('base.res_partner_12')
        product = self.env.ref('product.product_product_4_product_template')
        self.env['product.customerinfo'].create({
            'product_tmpl_id': product.id,
            'partner_id': partner.id,
            'product_code': '2323',
        })
        sale_order = self.env['sale.order'].create({
            'partner_id': partner.id,
        })
        products_customer_code = self.env['product.customerinfo'].search([
            ('product_code', '=', '2323'),
            ('partner_id', '=', partner.id),
        ])
        wizard = self.env['sale.order.lines_by_ref'].create({
            'sale_id': sale_order.id,
            'references': '2323',
        })
        wizard.action_simulate()
        self.assertEqual(len(wizard.line_ids), 1)
        wizard.action_create()
        self.assertEqual(len(sale_order.order_line), 0)
        self.assertEqual(len(products_customer_code), 1)
        self.assertEqual(len(products_customer_code[0].product_id), 0)
        self.assertEqual(len(
            products_customer_code[0].product_tmpl_id.product_variant_ids),
            products_customer_code[0].product_tmpl_id.product_variant_count)
        self.assertIn('1 variant for', wizard.line_ids[0].name)

    def test_sale_order_lines_customer_product_empty_product(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        partner = self.env.ref('base.res_partner_12')
        self.env['product.customerinfo'].create({
            'partner_id': partner.id,
            'product_code': '2323',
        })
        sale_order = self.env['sale.order'].create({
            'partner_id': partner.id,
        })
        products_customer_code = self.env['product.customerinfo'].search([
            ('product_code', '=', '2323'),
            ('partner_id', '=', partner.id),
        ])
        wizard = self.env['sale.order.lines_by_ref'].create({
            'sale_id': sale_order.id,
            'references': '2323',
        })
        wizard.action_simulate()
        wizard.action_create()
        self.assertEqual(len(sale_order.order_line), 0)
        self.assertEqual(len(products_customer_code), 1)
        self.assertEqual(len(products_customer_code[0].product_id), 0)
        self.assertEqual(len(products_customer_code[0].product_tmpl_id), 0)
        self.assertIn('not exists', wizard.line_ids[0].name)

    def test_sale_order_lines_permission_access(self):
        user = self.env.ref('base.user_demo')
        sale = self.env['sale.order'].create({
            'partner_id': self.partner.id,
        })
        wizard = self.env['sale.order.lines_by_ref'].with_user(
            user=user
        ).create({
            'sale_id': sale.id,
            'references': '2323',
        })
        wizard.action_simulate()

    def test_sale_order_partner_shipping_customerinfo(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        partner = self.env.ref('base.res_partner_12')
        partner_shipping = self.env['res.partner'].create({
            'name': 'Partner shipping test',
        })
        product = self.env['product.product'].create({
            'name': 'Product Test',
            'type': 'consu',
        })
        self.env['product.customerinfo'].create({
            'product_id': product.id,
            'partner_id': partner_shipping.id,
            'product_code': '2323',
        })
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
            'partner_shipping_id': partner_shipping.id,
        })
        customerinfo_obj = self.env['product.customerinfo']
        partner_customerinfo = customerinfo_obj.search([
            ('product_code', '=', '2323'),
            ('partner_id', '=', partner.id),
        ])
        partner_shipping_customerinfo = customerinfo_obj.search([
            ('product_code', '=', '2323'),
            ('name', '=', partner_shipping.id),
        ])
        self.assertEqual(sale.partner_id, partner)
        self.assertEqual(sale.partner_shipping_id, partner_shipping)
        wizard = self.env['sale.order.lines_by_ref'].create({
            'sale_id': sale.id,
            'references': '2323',
        })
        wizard.action_simulate()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_create()
        self.assertEqual(len(partner_customerinfo), 0)
        self.assertEqual(len(partner_shipping_customerinfo), 1)
        self.assertEqual(len(sale.order_line), 1)
        self.assertEqual(sale.order_line[0].product_id, product)

    def test_sale_partner_customerinfo_zero_price(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        fname = self.get_sample('sample_ok_customerinfo_1.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        partner = self.env.ref('base.res_partner_12')
        partner_shipping = self.env['res.partner'].create({
            'name': 'Partner shipping test',
        })
        customerinfo = self.env['product.customerinfo'].create({
            'product_id': self.product_1.id,
            'partner_id': partner.id,
            'product_code': 'R2323',
        })
        self.assertEqual(customerinfo.price, 0)
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
        })
        self.pricelist_item = self.env['product.pricelist.item'].create({
            'applied_on': '1_product',
            'base': 'partner',
            'name': 'Test pricelist item',
            'pricelist_id': sale.pricelist_id.id,
            'compute_price': 'fixed',
            'fixed_price': 6,
            'product_id': self.product_1.id,
        })
        customerinfo_obj = self.env['product.customerinfo']
        partner_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner.id),
        ])
        partner_shipping_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner_shipping.id),
        ])
        self.assertEqual(len(partner_customerinfo), 1)
        self.assertEqual(len(partner_shipping_customerinfo), 0)
        self.assertEqual(sale.partner_id, partner)
        self.assertEqual(sale.partner_shipping_id, partner)
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, 6)

    def test_sale_partner_customerinfo_price(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        fname = self.get_sample('sample_ok_customerinfo_1.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        partner = self.env.ref('base.res_partner_12')
        partner_shipping = self.env['res.partner'].create({
            'name': 'Partner shipping test',
        })
        customerinfo = self.env['product.customerinfo'].create({
            'product_id': self.product_1.id,
            'partner_id': partner.id,
            'product_code': 'R2323',
            'price': 4,
        })
        self.assertEqual(customerinfo.price, 4)
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
        })
        self.pricelist_item = self.env['product.pricelist.item'].create({
            'applied_on': '1_product',
            'base': 'partner',
            'name': 'Test pricelist item',
            'pricelist_id': sale.pricelist_id.id,
            'compute_price': 'fixed',
            'fixed_price': 6,
            'product_id': self.product_1.id,
        })
        customerinfo_obj = self.env['product.customerinfo']
        partner_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner.id),
        ])
        partner_shipping_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner_shipping.id),
        ])
        self.assertEqual(len(partner_customerinfo), 1)
        self.assertEqual(len(partner_shipping_customerinfo), 0)
        self.assertEqual(sale.partner_id, partner)
        self.assertEqual(sale.partner_shipping_id, partner)
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, 4)

    def test_sale_partner_shipping_customerinfo_price(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        fname = self.get_sample('sample_ok_customerinfo_1.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        partner = self.env.ref('base.res_partner_12')
        partner_shipping = self.env['res.partner'].create({
            'name': 'Partner shipping test',
        })
        customerinfo = self.env['product.customerinfo'].create({
            'product_id': self.product_1.id,
            'partner_id': partner_shipping.id,
            'product_code': 'R2323',
            'price': 4,
        })
        self.assertEqual(customerinfo.price, 4)
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
            'partner_shipping_id': partner_shipping.id,
        })
        self.pricelist_item = self.env['product.pricelist.item'].create({
            'applied_on': '1_product',
            'base': 'partner',
            'name': 'Test pricelist item',
            'pricelist_id': sale.pricelist_id.id,
            'compute_price': 'fixed',
            'fixed_price': 6,
            'product_id': self.product_1.id,
        })
        customerinfo_obj = self.env['product.customerinfo']
        partner_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner.id),
        ])
        partner_shipping_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner_shipping.id),
        ])
        self.assertEqual(len(partner_customerinfo), 0)
        self.assertEqual(len(partner_shipping_customerinfo), 1)
        self.assertEqual(sale.partner_id, partner)
        self.assertEqual(sale.partner_shipping_id, partner_shipping)
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, 4)

    def test_sale_partner_customerinfo_zero_price_pricelist_margin_01(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        if not self._pricelist_formula_margin():
            self.skipTest('Module Pricelist Formula Margin Not Installed')
        fname = self.get_sample('sample_ok_customerinfo_1.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        partner = self.env.ref('base.res_partner_12')
        partner_shipping = self.env['res.partner'].create({
            'name': 'Partner shipping test',
        })
        customerinfo = self.env['product.customerinfo'].create({
            'product_id': self.product_1.id,
            'partner_id': partner.id,
            'product_code': 'R2323',
        })
        self.assertEqual(customerinfo.price, 0)
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
        })
        self.pricelist_item = self.env['product.pricelist.item'].create({
            'applied_on': '3_global',
            'base': 'list_price',
            'name': 'Test pricelist item',
            'pricelist_id': sale.pricelist_id.id,
            'compute_price': 'margin',
            'percent_margin': 40,
        })
        customerinfo_obj = self.env['product.customerinfo']
        partner_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner.id),
        ])
        partner_shipping_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner_shipping.id),
        ])
        self.assertEqual(len(partner_customerinfo), 1)
        self.assertEqual(len(partner_shipping_customerinfo), 0)
        self.assertEqual(sale.partner_id, partner)
        self.assertEqual(sale.partner_shipping_id, partner)
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, 5)

    def test_sale_partner_customerinfo_zero_price_pricelist_margin_02(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        if not self._pricelist_formula_margin():
            self.skipTest('Module Pricelist Formula Margin Not Installed')
        fname = self.get_sample('sample_ok_customerinfo_1.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        partner = self.env.ref('base.res_partner_12')
        partner_shipping = self.env['res.partner'].create({
            'name': 'Partner shipping test',
        })
        customerinfo = self.env['product.customerinfo'].create({
            'product_id': self.product_1.id,
            'partner_id': partner.id,
            'product_code': 'R2323',
        })
        self.assertEqual(customerinfo.price, 0)
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
        })
        self.pricelist_item = self.env['product.pricelist.item'].create({
            'applied_on': '1_product',
            'base': 'list_price',
            'name': 'Test pricelist item',
            'pricelist_id': sale.pricelist_id.id,
            'compute_price': 'margin',
            'percent_margin': 60,
            'product_id': self.product_1.id,
        })
        customerinfo_obj = self.env['product.customerinfo']
        partner_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner.id),
        ])
        partner_shipping_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner_shipping.id),
        ])
        self.assertEqual(len(partner_customerinfo), 1)
        self.assertEqual(len(partner_shipping_customerinfo), 0)
        self.assertEqual(sale.partner_id, partner)
        self.assertEqual(sale.partner_shipping_id, partner)
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, 7.5)

    def test_partner_shipping_customerinfo_zero_price_pricelist_margin_03(
            self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        if not self._pricelist_formula_margin():
            self.skipTest('Module Pricelist Formula Margin Not Installed')
        fname = self.get_sample('sample_ok_customerinfo_1.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        partner = self.env.ref('base.res_partner_12')
        partner_shipping = self.env['res.partner'].create({
            'name': 'Partner shipping test',
        })
        customerinfo = self.env['product.customerinfo'].create({
            'product_id': self.product_1.id,
            'partner_id': partner_shipping.id,
            'product_code': 'R2323',
        })
        self.assertEqual(customerinfo.price, 0)
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
            'partner_shipping_id': partner_shipping.id,
        })
        self.pricelist_item = self.env['product.pricelist.item'].create({
            'applied_on': '3_global',
            'base': 'list_price',
            'name': 'Test pricelist item',
            'pricelist_id': sale.pricelist_id.id,
            'compute_price': 'margin',
            'percent_margin': 40,
        })
        customerinfo_obj = self.env['product.customerinfo']
        partner_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner.id),
        ])
        partner_shipping_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner_shipping.id),
        ])
        self.assertEqual(len(partner_customerinfo), 0)
        self.assertEqual(len(partner_shipping_customerinfo), 1)
        self.assertEqual(sale.partner_id, partner)
        self.assertEqual(sale.partner_shipping_id, partner_shipping)
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, 5)

    def test_partner_shipping_customerinfo_zero_price_pricelist_margin_04(
            self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        if not self._pricelist_formula_margin():
            self.skipTest('Module Pricelist Formula Margin Not Installed')
        fname = self.get_sample('sample_ok_customerinfo_1.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        partner = self.env.ref('base.res_partner_12')
        partner_shipping = self.env['res.partner'].create({
            'name': 'Partner shipping test',
        })
        customerinfo = self.env['product.customerinfo'].create({
            'product_id': self.product_1.id,
            'partner_id': partner_shipping.id,
            'product_code': 'R2323',
        })
        self.assertEqual(customerinfo.price, 0)
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
            'partner_shipping_id': partner_shipping.id,
        })
        self.pricelist_item = self.env['product.pricelist.item'].create({
            'applied_on': '1_product',
            'base': 'list_price',
            'name': 'Test pricelist item',
            'pricelist_id': sale.pricelist_id.id,
            'compute_price': 'margin',
            'percent_margin': 60,
            'product_id': self.product_1.id,
        })
        customerinfo_obj = self.env['product.customerinfo']
        partner_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner.id),
        ])
        partner_shipping_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner_shipping.id),
        ])
        self.assertEqual(len(partner_customerinfo), 0)
        self.assertEqual(len(partner_shipping_customerinfo), 1)
        self.assertEqual(sale.partner_id, partner)
        self.assertEqual(sale.partner_shipping_id, partner_shipping)
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, 7.5)

    def test_sale_partner_shipping_customerinfo_zero_price(self):
        if not self._customer_info():
            self.skipTest('Module SupplierInfo Not Installed')
        fname = self.get_sample('sample_ok_customerinfo_1.xlsx')
        file = base64.b64encode(open(fname, 'rb').read())
        partner = self.env.ref('base.res_partner_12')
        partner_shipping = self.env['res.partner'].create({
            'name': 'Partner shipping test',
        })
        customerinfo = self.env['product.customerinfo'].create({
            'product_id': self.product_1.id,
            'partner_id': partner_shipping.id,
            'product_code': 'R2323',
        })
        self.assertEqual(customerinfo.price, 0)
        sale = self.env['sale.order'].create({
            'partner_id': partner.id,
            'partner_shipping_id': partner_shipping.id,
        })
        self.pricelist_item = self.env['product.pricelist.item'].create({
            'applied_on': '1_product',
            'base': 'partner',
            'name': 'Test pricelist item',
            'pricelist_id': sale.pricelist_id.id,
            'compute_price': 'fixed',
            'fixed_price': 6,
            'product_id': self.product_1.id,
        })
        customerinfo_obj = self.env['product.customerinfo']
        partner_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner.id),
        ])
        partner_shipping_customerinfo = customerinfo_obj.search([
            ('product_code', '=', 'R2323'),
            ('partner_id', '=', partner_shipping.id),
        ])
        self.assertEqual(len(partner_customerinfo), 0)
        self.assertEqual(len(partner_shipping_customerinfo), 1)
        self.assertEqual(sale.partner_id, partner)
        self.assertEqual(sale.partner_shipping_id, partner_shipping)
        wizard = self.env['sale.order.lines_by_ref.xlsx'].create({
            'sale_id': sale.id,
            'xlsx_file': file,
        })
        wizard.action_simulate_import_lines_from_xlsx()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.action_import_lines_from_xlsx()
        self.assertEqual(len(sale.order_line), 1)
        line = sale.order_line[0]
        self.assertEqual(line.product_id, self.product_1)
        self.assertEqual(line.product_uom_qty, 1)
        self.assertEqual(line.price_unit, 6)
