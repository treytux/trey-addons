###############################################################################
# For copyright and license notices, see __manifest__.py file in root directory
###############################################################################
import base64
import os

import openpyxl
from dateutil.relativedelta import relativedelta
from odoo import fields
from odoo.tests import common


class TestContractLineLotImport(common.TransactionCase):

    def setUp(self):
        super().setUp()
        self.partner = self.env['res.partner'].create({
            'name': 'Test partner',
            'is_company': True,
        })
        self.product_01 = self.env['product.product'].create({
            'type': 'product',
            'company_id': False,
            'name': 'Test product 1',
            'standard_price': 10,
            'default_code': '01-PROD',
            'list_price': 100,
        })
        self.product_02 = self.env['product.product'].create({
            'type': 'product',
            'company_id': False,
            'name': 'Test product 2',
            'standard_price': 10,
            'default_code': '02-PROD',
            'list_price': 50,
        })
        self.lot_01 = self.env['stock.production.lot'].create({
            'name': '987654321',
            'product_id': self.product_02.id,
        })
        self.contract = self.env['contract.contract'].create({
            'name': 'CTRTEST',
            'partner_id': self.partner.id,
        })
        self.assertEqual(len(self.contract.contract_line_ids), 0)
        self.contract_line = self.env['contract.line'].create({
            'name': 'Contract line test',
            'product_id': self.product_01.id,
            'contract_id': self.contract.id,
            'quantity': 1,
            'price_unit': 50,
            'recurring_rule_type': 'quarterly',
            'date_start': fields.Date.today(),
        })
        self.xlsx_name = 'test_import_lot_contract_line.xlsx'

    def create_excel(self, path, partners):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet['A1'] = 'Suscripción'
        sheet['B1'] = 'Modalidad Mantenimiento'
        sheet['C1'] = 'S/N'
        sheet['D1'] = 'P/N'
        cont = 2
        for _index in partners:
            sheet[f'A{cont}'] = self.contract.name
            sheet[f'B{cont}'] = self.product_01.default_code
            sheet[f'C{cont}'] = self.lot_01.name
            sheet[f'D{cont}'] = self.product_02.default_code
            cont += 1
        workbook.save(path)
        return workbook

    def create_excel_alternative(self, path, partners, product_lot):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet['A1'] = 'Suscripción'
        sheet['B1'] = 'Modalidad Mantenimiento'
        sheet['C1'] = 'S/N'
        sheet['D1'] = 'P/N'
        sheet['A2'] = self.contract.name
        sheet['B2'] = self.product_01.default_code
        sheet['C2'] = ''
        sheet['D2'] = product_lot.default_code
        workbook.save(path)
        return workbook

    def create_excel_no_product(self, path, partners, lot):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet['A1'] = 'Suscripción'
        sheet['B1'] = 'Modalidad Mantenimiento'
        sheet['C1'] = 'S/N'
        sheet['D1'] = 'P/N'
        sheet['A2'] = self.contract.name
        sheet['B2'] = self.product_01.default_code
        sheet['C2'] = lot
        sheet['D2'] = ''
        workbook.save(path)
        return workbook

    def check_xlsx_exists(self):
        path = os.path.dirname(__file__)
        path = os.path.join(path, self.xlsx_name)
        if os.path.isfile(path):
            os.remove(path)
        return path

    def create_wizard(self, xlsx_path):
        file = base64.b64encode(open(xlsx_path, 'rb').read())
        return self.env['contract.line.lot.import'].create({
            'file': file,
        })

    def test_import_lot_contract_line(self):
        xlsx_path = self.check_xlsx_exists()
        self.create_excel(xlsx_path, [self.partner])
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        wizard.action_simulate()
        self.assertEquals(len(wizard.line_ids), 0)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        wizard.button_accept()
        self.assertEqual(len(self.contract_line.lot_ids), 1)

    def test_error_no_contract_find(self):
        xlsx_path = self.check_xlsx_exists()
        self.create_excel(xlsx_path, [self.partner])
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        self.contract.unlink()
        wizard.action_simulate()
        self.assertEquals(len(wizard.line_ids), 1)
        self.assertIn(
            'No contract found', wizard.line_ids[0].name)

    def test_error_multiple_contract_find(self):
        self.env['contract.contract'].create({
            'name': 'CTRTEST',
            'partner_id': self.partner.id,
        })
        xlsx_path = self.check_xlsx_exists()
        self.create_excel(xlsx_path, [self.partner])
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        wizard.action_simulate()
        self.assertEquals(len(wizard.line_ids), 1)
        self.assertIn(
            'More than one contract', wizard.line_ids[0].name)

    def test_error_no_product_contract(self):
        xlsx_path = self.check_xlsx_exists()
        self.create_excel(xlsx_path, [self.partner])
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        self.product_01.unlink()
        wizard.action_simulate()
        self.assertEquals(len(wizard.line_ids), 1)
        self.assertIn(
            'No product', wizard.line_ids[0].name)

    def test_error_multiple_product_contract_find(self):
        xlsx_path = self.check_xlsx_exists()
        self.create_excel(xlsx_path, [self.partner])
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        product_copy = self.product_01.copy()
        product_copy.name = 'Test product 1 copy'
        self.assertEqual(
            product_copy.default_code, self.product_01.default_code)
        wizard.action_simulate()
        self.assertEquals(len(wizard.line_ids), 1)
        self.assertIn(
            'More than one product', wizard.line_ids[0].name)

    def test_error_no_product_line_contract(self):
        self.assertEqual(len(self.contract.contract_line_ids), 1)
        self.contract_line.cancel()
        xlsx_path = self.check_xlsx_exists()
        self.create_excel(xlsx_path, [self.partner])
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        self.contract_line.unlink()
        self.assertEqual(len(self.contract.contract_line_ids), 0)
        wizard.action_simulate()
        self.assertEquals(len(wizard.line_ids), 1)
        self.assertIn(
            'No contract line found with this product', wizard.line_ids[0].name)

    def test_error_not_lot_found(self):
        self.assertEqual(len(self.contract.contract_line_ids), 1)
        xlsx_path = self.check_xlsx_exists()
        self.create_excel(xlsx_path, [self.partner])
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        self.lot_01.unlink()
        wizard.action_simulate()
        self.assertEquals(len(wizard.line_ids), 1)
        self.assertIn(
            'New lot created', wizard.line_ids[0].name)

    def test_error_multiple_product_lines_contract(self):
        self.assertEqual(len(self.contract.contract_line_ids), 1)
        self.env['contract.line'].create({
            'name': 'Contract line test',
            'product_id': self.product_01.id,
            'contract_id': self.contract.id,
            'quantity': 1,
            'price_unit': 80,
            'recurring_rule_type': 'quarterly',
            'date_start': fields.Date.today() + relativedelta(days=10),
        })
        self.assertEqual(len(self.contract.contract_line_ids), 2)
        xlsx_path = self.check_xlsx_exists()
        self.create_excel(xlsx_path, [self.partner])
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        wizard.action_simulate()
        self.assertEquals(len(wizard.line_ids), 0)
        wizard.button_accept()
        self.assertEqual(len(self.contract_line.lot_ids), 1)

    def test_column_lot_ok_product_column_empty(self):
        xlsx_path = self.check_xlsx_exists()
        self.create_excel_no_product(xlsx_path, [self.partner], 'NOLOT')
        wizard = self.create_wizard(xlsx_path)
        wizard.action_simulate()
        self.assertEqual(len(wizard.line_ids), 1)
        self.assertIn(
            'The lot cannot be created because the product is not defined',
            wizard.line_ids[0].name)

    def test_create_lot_ok(self):
        xlsx_path = self.check_xlsx_exists()
        self.create_excel(xlsx_path, [self.partner])
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        lot_name = self.lot_01.name
        self.lot_01.unlink()
        lots = self.env['stock.production.lot'].search([
            ('name', '=', lot_name),
        ])
        self.assertEqual(len(lots), 0)
        wizard.action_simulate()
        self.assertEquals(len(wizard.line_ids), 1)
        self.assertIn('New lot created', wizard.line_ids[0].name)
        lots = self.env['stock.production.lot'].search([
            ('name', '=', lot_name),
        ])
        self.assertEqual(len(lots), 0)
        wizard.button_accept()
        lots = self.env['stock.production.lot'].search([
            ('name', '=', lot_name),
        ])
        self.assertEqual(len(lots), 1)
        self.assertEqual(len(self.contract_line.lot_ids), 1)
        self.assertEqual(self.contract_line.lot_ids[0], lots[0])

    def test_create_lot_no_products(self):
        xlsx_path = self.check_xlsx_exists()
        self.create_excel(xlsx_path, [self.partner])
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        lot_name = self.lot_01.name
        self.lot_01.unlink()
        lots = self.env['stock.production.lot'].search([
            ('name', '=', lot_name),
        ])
        self.assertEqual(len(lots), 0)
        self.env['product.product'].create({
            'type': 'product',
            'company_id': False,
            'name': 'Test product copy 3',
            'standard_price': 10,
            'default_code': '02-PROD',
            'list_price': 50,
        })
        wizard.action_simulate()
        self.assertEquals(len(wizard.line_ids), 1)
        self.assertIn(
            'Row 1: More than one product', wizard.line_ids[0].name)

    def test_create_lot_create_product_import(self):
        xlsx_path = self.check_xlsx_exists()
        self.create_excel(xlsx_path, [self.partner])
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        lot_name = self.lot_01.name
        self.lot_01.unlink()
        lots = self.env['stock.production.lot'].search([
            ('name', '=', lot_name),
        ])
        self.assertEqual(len(lots), 0)
        product_code = self.product_02.default_code
        self.product_02.unlink()
        products = self.env['product.product'].search([
            ('default_code', '=', product_code),
        ])
        self.assertEqual(len(products), 0)
        wizard.action_simulate()
        self.assertEquals(len(wizard.line_ids), 2)
        self.assertIn('New product created', wizard.line_ids[0].name)
        self.assertIn('New lot created', wizard.line_ids[1].name)
        lots = self.env['stock.production.lot'].search([
            ('name', '=', lot_name),
        ])
        self.assertEqual(len(lots), 0)
        products = self.env['product.product'].search([
            ('default_code', '=', product_code),
        ])
        self.assertEqual(len(products), 0)
        wizard.button_accept()
        lots = self.env['stock.production.lot'].search([
            ('name', '=', lot_name),
        ])
        self.assertEqual(len(lots), 1)
        products = self.env['product.product'].search([
            ('default_code', '=', product_code),
        ])
        self.assertEqual(len(products), 1)
        self.assertEqual(len(self.contract_line.lot_ids), 1)
        self.assertEqual(self.contract_line.lot_ids[0], lots[0])

    def test_no_lot_column_lot_empty(self):
        xlsx_path = self.check_xlsx_exists()
        product_test = self.env['product.product'].create({
            'type': 'product',
            'company_id': False,
            'name': 'Test product no lot column empty',
            'standard_price': 10,
            'default_code': 'PRD-TEST',
            'list_price': 100,
        })
        self.create_excel_alternative(
            xlsx_path, [self.partner], product_test)
        product_test.unlink()
        products = self.env['product.product'].search([
            ('default_code', '=', 'PRD-TEST'),
        ])
        self.assertEqual(len(products), 0)
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        wizard.action_simulate()
        self.assertEquals(len(wizard.line_ids), 1)
        self.assertIn(
            'The lot cannot be assigned because the S/N column is not set',
            wizard.line_ids[0].name)

    def test_cannot_select_lot_column_empty(self):
        xlsx_path = self.check_xlsx_exists()
        product_test = self.env['product.product'].create({
            'type': 'product',
            'company_id': False,
            'name': 'Test product no lot column empty',
            'standard_price': 10,
            'default_code': 'PRD-TEST',
            'list_price': 100,
        })
        self.create_excel_alternative(
            xlsx_path, [self.partner], product_test)
        products = self.env['product.product'].search([
            ('default_code', '=', 'PRD-TEST'),
        ])
        self.assertEqual(len(products), 1)
        lots = self.env['stock.production.lot'].search([
            ('product_id', '=', product_test.id),
        ])
        self.assertEqual(len(lots), 0)
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        wizard.action_simulate()
        self.assertEquals(len(wizard.line_ids), 1)
        self.assertIn(
            'The lot cannot be assigned because the S/N column is not set',
            wizard.line_ids[0].name)

    def test_check_alternative_columns(self):
        xlsx_path = self.check_xlsx_exists()
        self.create_excel_alternative(
            xlsx_path, [self.partner], self.product_02)
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        wizard.action_simulate()
        self.assertEquals(len(wizard.line_ids), 0)
        wizard.button_accept()
        self.assertEqual(len(self.contract_line.lot_ids), 1)

    def test_check_alternative_columns_error(self):
        xlsx_path = self.check_xlsx_exists()
        product_03 = self.env['product.product'].create({
            'type': 'product',
            'company_id': False,
            'name': 'Test product 3',
            'standard_price': 20,
            'default_code': '03-PROD',
            'list_price': 86,
        })
        self.create_excel_alternative(xlsx_path, [self.partner], product_03)
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        wizard.action_simulate()
        self.assertEquals(len(wizard.line_ids), 1)
        self.assertIn(
            'The lot cannot be assigned because the S/N column is not set',
            wizard.line_ids[0].name)

    def test_repeat_import_lot_contract_line(self):
        xlsx_path = self.check_xlsx_exists()
        self.create_excel(xlsx_path, [self.partner])
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        wizard.action_simulate()
        self.assertEqual(len(wizard.line_ids), 0)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        wizard.button_accept()
        self.assertEqual(len(self.contract_line.lot_ids), 1)
        wizard_02 = self.create_wizard(xlsx_path)
        self.assertTrue(wizard_02.file)
        self.assertEqual(len(self.contract_line.lot_ids), 1)
        self.assertEqual(self.contract_line.lot_ids[0], self.lot_01)
        wizard_02.action_simulate()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard_02.button_accept()
        self.assertEqual(len(self.contract_line.lot_ids), 1)
        self.assertEqual(self.contract_line.lot_ids[0], self.lot_01)

    def test_import_lot_two_contracts(self):
        xlsx_path = self.check_xlsx_exists()
        self.create_excel(xlsx_path, [self.partner])
        wizard = self.create_wizard(xlsx_path)
        self.assertTrue(wizard.file)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        wizard.action_simulate()
        self.assertEqual(len(wizard.line_ids), 0)
        wizard.button_accept()
        self.assertEqual(len(self.contract_line.lot_ids), 1)
        self.assertEqual(self.contract_line.lot_ids[0], self.lot_01)
        contract_02 = self.env['contract.contract'].create({
            'name': 'Contract example 2',
            'partner_id': self.partner.id,
        })
        self.assertEqual(len(contract_02.contract_line_ids), 0)
        contract_line_02 = self.env['contract.line'].create({
            'name': 'Contract line test',
            'product_id': self.product_01.id,
            'contract_id': contract_02.id,
            'quantity': 1,
            'price_unit': 80,
            'recurring_rule_type': 'quarterly',
            'date_start': fields.Date.today(),
        })
        workbook = self.create_excel(xlsx_path, [self.partner])
        sheet = workbook.active
        sheet['A2'].value = contract_02.name
        self.assertEqual(sheet['A2'].value, contract_02.name)
        workbook.save(xlsx_path)
        wizard_02 = self.create_wizard(xlsx_path)
        self.assertTrue(wizard_02.file)
        wizard_02.action_simulate()
        self.assertEqual(len(wizard_02.line_ids), 0)
        self.assertEqual(len(self.contract_line.lot_ids), 1)
        self.assertEqual(len(contract_line_02.lot_ids), 0)
        wizard_02.button_accept()
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        self.assertEqual(len(contract_line_02.lot_ids), 1)
        self.assertEqual(contract_line_02.lot_ids[0], self.lot_01)

    def test_import_lot_contract_line_search_by_code(self):
        self.assertFalse(self.contract.code)
        self.contract.code = 'REFX01'
        self.assertEqual(self.contract.code , 'REFX01')
        xlsx_path = self.check_xlsx_exists()
        workbook = self.create_excel(xlsx_path, [self.partner])
        sheet = workbook.active
        sheet['A2'].value = self.contract.code
        self.assertEqual(sheet['A2'].value, self.contract.code)
        wizard = self.create_wizard(xlsx_path)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        wizard.action_simulate()
        self.assertEqual(len(wizard.line_ids), 0)
        self.assertEqual(len(self.contract_line.lot_ids), 0)
        wizard.button_accept()
        self.assertEqual(len(self.contract_line.lot_ids), 1)
