###############################################################################
# For copyright and license notices, see __manifest__.py file in root directory
###############################################################################
import base64
import unicodedata
from datetime import datetime
from io import BytesIO

import pandas as pd
from odoo import _, fields, models
from odoo.exceptions import UserError
from openpyxl import load_workbook


class ProductVariantConverterWizard(models.TransientModel):
    _name = 'product.variant.converter.wizard'
    _description = 'Product Variant Converter Wizard'

    supplier = fields.Selection([
        ('brooks', 'Brooks'),
        ('salomon', 'Salomon'),
        ('asics', 'Asics'),
        ('new_balance', 'New Balance'),
        ('hoka', 'Hoka'),
        ('mizuno', 'Mizuno'),
    ],
        string='Supplier',
        required=True,
    )
    input_file = fields.Binary(
        string='File to Convert',
        required=True,
    )
    input_filename = fields.Char(
        string='Input Filename',
    )
    output_file = fields.Binary(
        string='Converted File',
        readonly=True,
    )
    output_filename = fields.Char(
        string='Output Filename',
        readonly=True,
    )
    state = fields.Selection(
        string='State',
        selection=[
            ('upload_file', 'Upload File to Convert'),
            ('download_file', 'Download Converted File'),
        ],
        required=True,
        default='upload_file',
    )

    def _check_non_empty_file(self, df):
        if df.empty:
            raise UserError(_('File is empty.'))
        return df

    def _check_duplicate_columns(self):
        stream = BytesIO(base64.b64decode(self.input_file))
        wb = load_workbook(stream, read_only=True)
        if self.supplier == 'brooks':
            sheets_to_check = ['Pricat Completo']
        elif self.supplier == 'salomon':
            sheets_to_check = ['Worksheet']
        elif self.supplier == 'asics':
            sheets_to_check = ['Sheet1']
        elif self.supplier == 'new_balance':
            sheets_to_check = ['Pedido Elastic']
        elif self.supplier == 'hoka':
            sheets_to_check = ['K25']
        elif self.supplier == 'mizuno':
            sheets_to_check = [
                'Calzado',
                'Apparel & Accesorios',
                'Accesorios individual',
            ]
        else:
            sheets_to_check = []
        all_duplicates = {}
        for sheet_name in sheets_to_check:
            if sheet_name not in wb.sheetnames:
                raise UserError(
                    _('Error, sheet "%s" not found in the file.') % sheet_name)
            sheet = wb[sheet_name]
            headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
            duplicates = [h for h in set(headers) if headers.count(h) > 1 and h]
            if duplicates:
                all_duplicates[sheet_name] = duplicates
        if all_duplicates:
            errors = [
                f"'{sheet}': {', '.join(dups)}"
                for sheet, dups in all_duplicates.items()
            ]
            raise UserError(
                _('Error, duplicate columns in sheets: %s') % "; ".join(errors)
            )
        stream.seek(0)

    def _validate_required_columns(self, df, required_columns, sep):
        missing = [col for col in required_columns if col not in df.columns]
        if missing:
            raise UserError(
                _('File must contain columns: %s') % sep.join(missing))

    def _validate_non_empty_rows(self, df, columns, header_row=0):
        for col in columns:
            if col not in df.columns:
                continue
            empty_rows = df[col].isnull() | (df[col] == '')
            if empty_rows.any():
                if self.supplier == 'brooks':
                    row_numbers = (
                        df[empty_rows].index + header_row + 3
                    ).tolist()
                else:
                    row_numbers = (
                        df[empty_rows].index + header_row + 2
                    ).tolist()
                row_str = ', '.join(str(n) for n in row_numbers)
                raise UserError(_(
                    'Column \'%s\' is empty in rows: %s') % (col, row_str))

    def _validate_file_type(self):
        try:
            stream = BytesIO(base64.b64decode(self.input_file))
            xl = pd.ExcelFile(stream)
        except Exception:
            raise UserError(_(
                'Invalid file type. Please upload an Excel file.'))
        return xl

    def _check_sheets(self, xl):
        if not xl.sheet_names:
            raise UserError(_('File has no sheets.'))
        all_sheets_empty = all(
            xl.parse(sheet).dropna(how='all').empty
            for sheet in xl.sheet_names
        )
        if all_sheets_empty:
            raise UserError(_('File sheets are empty.'))

    def _xl_parser(self, xl):
        if self.supplier == 'brooks':
            df = xl.parse(
                'Pricat Completo',
                header=1,
                na_values=['NULL'],
                keep_default_na=None,
                converters={
                    'EAN': str,
                    'Precio Coste': str,
                },
            )
        elif self.supplier == 'salomon':
            df = xl.parse(
                'Worksheet',
                header=0,
                na_values=['NULL'],
                keep_default_na=None,
                converters={
                    'EAN': str,
                },
            )
        elif self.supplier == 'asics':
            df = xl.parse(
                'Sheet1',
                header=0,
                na_values=['NULL'],
                keep_default_na=None,
                converters={
                    'EAN': str,
                },
            )
        elif self.supplier == 'new_balance':
            df = xl.parse(
                'Pedido Elastic',
                header=0,
                na_values=['NULL'],
                keep_default_na=None,
                converters={
                    'UPC': str,
                },
            )
        elif self.supplier == 'hoka':
            df = xl.parse(
                'K25',
                header=0,
                na_values=['NULL'],
                keep_default_na=None,
                converters={
                    'REFERENCIA': str,
                    'EAN': str,
                },
            )
        elif self.supplier == 'mizuno':
            sheet_names = [
                'Calzado', 'Apparel & Accesorios', 'Accesorios individual']
            all_data = []
            for sheet in sheet_names:
                df_sheet = xl.parse(
                    sheet,
                    header=0,
                    na_values=['NULL'],
                    keep_default_na=None,
                    converters={
                        'Código Artículo': str,
                        'EAN': str,
                        'Talla': str,
                        'Talla Eur': str,
                    },
                )
                if sheet == 'Calzado':
                    df_sheet.rename(columns={
                        'Talla Eur': 'Talla',
                    }, inplace=True)
                all_data.append(df_sheet)
            df = pd.concat(all_data, ignore_index=True)
        df = df.dropna(how='all')
        df = df.dropna(how='all', axis=1)
        return df

    def _get_required_columns(self):
        if self.supplier == 'brooks':
            required_columns = ['EAN', 'Style/Color']
        elif self.supplier == 'salomon':
            required_columns = ['Color', 'EAN', 'Código SAP']
        elif self.supplier == 'asics':
            required_columns = ['ZZTRADING_CODE', 'EAN']
        elif self.supplier == 'new_balance':
            required_columns = ['Style Number', 'UPC', 'Precio de venta']
        elif self.supplier == 'hoka':
            required_columns = ['REFERENCIA', 'EAN']
        elif self.supplier == 'mizuno':
            required_columns = ['Código Artículo', 'EAN']
        return required_columns

    def _files_cleaning(self, df):
        df.columns = df.columns.str.strip()
        df.columns = [
            unicodedata.normalize('NFKC', str(col))
            .replace('\u00A0', '')
            .replace('\t', '')
            .strip()
            for col in df.columns
        ]
        df.replace({r'[\u00A0\t]+': ''}, regex=True, inplace=True)
        df = df.where(pd.notnull(df), None)
        return df

    def _read_file(self, xl, sep=', '):
        if self.supplier == 'brooks':
            if 'Pricelist' not in xl.sheet_names:
                raise UserError(_(
                    'File must contain a sheet named \'Pricelist\'.'))
            if 'Pricat Completo' not in xl.sheet_names:
                raise UserError(_(
                    'File must contain a sheet named \'Pricat Completo\'.'))
        if self.supplier == 'salomon':
            if 'Worksheet' not in xl.sheet_names:
                raise UserError(_(
                    'File must contain a sheet named \'Worksheet\'.'))
        if self.supplier == 'asics':
            if 'Sheet1' not in xl.sheet_names:
                raise UserError(_(
                    'File must contain a sheet named \'Sheet1\'.'))
        if self.supplier == 'new_balance':
            if 'Pedido Elastic' not in xl.sheet_names:
                raise UserError(_(
                    'File must contain a sheet named \'Pedido Elastic\'.'))
        if self.supplier == 'hoka':
            if 'K25' not in xl.sheet_names:
                raise UserError(_(
                    'File must contain a sheet named \'K25\'.'))
        if self.supplier == 'mizuno':
            required_sheets = [
                'Calzado',
                'Apparel & Accesorios',
                'Accesorios individual',
            ]
            for sheet in required_sheets:
                if sheet not in xl.sheet_names:
                    raise UserError(_(
                        'File must contain a sheet named \'%s\'.') % sheet)
        df = self._xl_parser(xl)
        df = self._files_cleaning(df)
        self._check_duplicate_columns()
        required_columns = self._get_required_columns()
        self._validate_required_columns(df, required_columns, sep)
        self._validate_non_empty_rows(df, required_columns, header_row=0)
        if df.empty:
            raise UserError(_('File is empty.'))
        return df

    def _get_default_values(self):
        return {
            'categ_id': 'All',
            'type': 'product',
            '*TMPL*sale_ok': True,
            '*TMPL*purchase_ok': True,
            '*TMPL*invoice_policy': 'delivery',
            'uom_id': _('Units'),
            'uom_po_id': _('Units'),
        }

    def _convert_brooks_file(self, df):
        def get_name(row):
            if row.get('Gender') == 'Womens':
                return _('%s Woman') % row.get('Item Description', '')
            return row.get('Item Description', '')
        standard_price = self.parse_price(df['Precio Coste'])
        list_price = self.parse_price(df['Precio PVP'])
        df_converted = {
            'product_tmpl_code': df['Style/Color'],
            'default_code': df['Style/Color'],
            'attribute: Color': df['Color'],
            'attribute: Size': df['Size EU'],
            'barcode': df['EAN'],
            'name': df.apply(get_name, axis=1),
            'standard_price': standard_price,
            '*TMPL*list_price': list_price,
            'weight': df['Gross wt (kg)'],
        }
        default_values = self._get_default_values()
        df_converted.update(default_values)
        return pd.DataFrame(df_converted)

    def _convert_salomon_file(self, df):
        def get_name(row):
            if row.get('Género') == 'Mujer':
                return _('%s Woman') % row.get('Descripción', '')
            return row.get('Descripción', '')
        standard_price = self.parse_price(df['Neto'])
        list_price = self.parse_price(df['SRP'])
        df_converted = {
            'product_tmpl_code': df['Código SAP'],
            'default_code': df['Código SAP'],
            'attribute: Color': df['Color'],
            'attribute: Size': df['Talla'],
            'barcode': df['EAN'],
            'name': df.apply(get_name, axis=1),
            'standard_price': standard_price,
            '*TMPL*list_price': list_price,
        }
        default_values = self._get_default_values()
        df_converted.update(default_values)
        return pd.DataFrame(df_converted)

    def _convert_asics_file(self, df):
        def get_name(row):
            if row.get('ZZGENDER_TXT') == 'Women':
                return _('%s Woman') % row.get('MAKTX', '')
            return row.get('MAKTX', '')
        standard_price = self.parse_price(df['UNIT_PRICE'])
        list_price = self.parse_price(df['RETAIL_PRICE'])
        df_converted = {
            'product_tmpl_code': df['ZZTRADING_CODE'],
            'default_code': df['ZZTRADING_CODE'],
            'attribute: Color': df['ZZCOMM_COLOR'],
            'attribute: Size': df['SIZE'],
            'barcode': df['EAN'],
            'name': df.apply(get_name, axis=1),
            'standard_price': standard_price,
            '*TMPL*list_price': list_price,
            'weight': df['GROSSWEIGHT'],
        }
        default_values = self._get_default_values()
        df_converted.update(default_values)
        return pd.DataFrame(df_converted)

    def _convert_new_balance_file(self, df):
        def get_name(row):
            if row.get('Gender') == "Women's":
                return _('%s Woman') % row.get('Style Name', '')
            return row.get('Style Name', '')
        standard_price = self.parse_price(df['Wholesale Price'])
        list_price = self.parse_price(df['Precio de venta'])
        df_converted = {
            'product_tmpl_code': df['Style Number'],
            'default_code': df['Style Number'],
            'attribute: Color': df['Color Name'],
            'attribute: Size': df['Size'],
            'barcode': '0' + df['UPC'],
            'name': df.apply(get_name, axis=1),
            'standard_price': standard_price,
            '*TMPL*list_price': list_price,
        }
        default_values = self._get_default_values()
        df_converted.update(default_values)
        return pd.DataFrame(df_converted)

    def _convert_hoka_file(self, df):
        def get_name(row):
            if row.get('GÉNERO') == 'Mujer':
                return _('%s Woman') % row.get('DESCRIPCIÓN', '')
            return row.get('DESCRIPCIÓN', '')
        standard_price = self.parse_price(df['Neto'])
        list_price = self.parse_price(df['PVPR'])
        df_converted = {
            'product_tmpl_code': df['REFERENCIA'],
            'default_code': df['REFERENCIA'],
            'attribute: Color': df['COLOR'],
            'attribute: Size': df['TALLA'].str.strip(),
            'barcode': df['EAN'],
            'name': df.apply(get_name, axis=1),
            'standard_price': standard_price,
            '*TMPL*list_price': list_price,
        }
        default_values = self._get_default_values()
        df_converted.update(default_values)
        return pd.DataFrame(df_converted)

    def _convert_mizuno_file(self, df):
        def get_name(row):
            if row.get('Género') == 'Mujer':
                return _('%s Woman') % row.get('Descripción', '')
            return row.get('Descripción', '')
        standard_price = self.parse_price(df['PRECIO TARIFA'])
        list_price = self.parse_price(df['P.V.P'])
        df_converted = {
            'product_tmpl_code': df['Código Artículo'],
            'default_code': df['Código Artículo'],
            'attribute: Color': df['Descripción Color'],
            'attribute: Size': df['Talla'].str.strip(),
            'barcode': df['EAN'],
            'name': df.apply(get_name, axis=1),
            'standard_price': standard_price,
            '*TMPL*list_price': list_price,
        }
        default_values = self._get_default_values()
        df_converted.update(default_values)
        return pd.DataFrame(df_converted)

    def parse_price(self, price):
        try:
            price = (
                price.astype(str).str.replace('€', '').str.strip().str.replace(
                    ',', '.').astype(float))
            return price
        except Exception as e:
            invalid_rows = (
                pd.to_numeric(
                    price.astype(str).str.replace('€', '').str.strip(
                    ).str.replace(',', '.'),
                    errors='coerce'
                ).isna()
            )
            row_numbers = (price[invalid_rows].index + 2).tolist()
            row_str = ', '.join(str(n) for n in row_numbers)
            if 'could not convert string to float: ' in e.args[0]:
                e = _(
                    'Error parsing price. Please check the format. '
                    'Ensure that the price is a valid number.')
                raise UserError(
                    _('%s Column \'%s\', rows: %s.') % (
                        e, price.name, row_str
                    )
                )

            raise UserError(_('Error: %s.') % e)

    def _validate_files(self):
        xl = self._validate_file_type()
        self._check_sheets(xl)
        return xl

    def _conversion(self, xl):
        df = self._read_file(xl)
        if self.supplier == 'brooks':
            data = self._convert_brooks_file(df)
        elif self.supplier == 'salomon':
            data = self._convert_salomon_file(df)
        elif self.supplier == 'asics':
            data = self._convert_asics_file(df)
        elif self.supplier == 'new_balance':
            data = self._convert_new_balance_file(df)
        elif self.supplier == 'hoka':
            data = self._convert_hoka_file(df)
        elif self.supplier == 'mizuno':
            data = self._convert_mizuno_file(df)
        return data

    def action_accept(self):
        self.ensure_one()
        xl = self._validate_files()
        data = self._conversion(xl)
        output_stream = BytesIO()
        data.to_excel(output_stream, index=False, engine='openpyxl')
        output_stream.seek(0)
        self.output_file = base64.b64encode(output_stream.read())
        current_date = datetime.now().strftime('%Y-%m-%d')
        self.output_filename = (
            f"{self.supplier}_converted_template_{current_date}.xlsx"
        )
        self.state = 'download_file'
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'product.variant.converter.wizard',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }
